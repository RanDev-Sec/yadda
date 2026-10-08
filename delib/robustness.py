"""robustness - runs MITRE's Detection Coverage Calculator (Summiting the Pyramid) on an environment's enabled
YARA-L rules, through shadow Sigma files built from each rule's events: section."""
from __future__ import annotations

import hashlib

import csv
import re
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path
from delib.config import _rmtree, envdir, die, read_csv, SHARED, write_csv, write_text, yaml_dump_plain
from delib.yaral import _tree_map, _tree_preds, _values, _yl_events, _yl_events_tree
from delib.upstream import _dcc, update_dcc
from delib.telemetry import _mapping, _row_dcs
from delib.facts import _rule_facts


# ---------------------------------------------------------------- robustness: YARA-L -> shadow Sigma -> MITRE DCC
# The calculator only reads Sigma. A "shadow Sigma" file carries what it scores - ATT&CK tags, a logsource, the fields
# each event variable tests and how they combine - and is never meant to run. Field names are not hard-coded: each UDM
# field is translated to OCSF (shared/udm_to_ocsf.csv), then to whatever Sigma field the *installed* calculator
# lists for that OCSF field in its own scoring dictionary, so fields renamed in a newer calculator are picked up.
ENDPOINT_CATEGORIES = {  # UDM event type -> Sigma logsource category (checked against the calculator's dictionary)
    "PROCESS_LAUNCH": "process_creation", "NETWORK_CONNECTION": "network_connection", "NETWORK_DNS": "dns_query",
    "FILE_CREATION": "file_event", "FILE_MODIFICATION": "file_change", "FILE_DELETION": "file_delete",
    "REGISTRY_CREATION": "registry_add", "REGISTRY_MODIFICATION": "registry_set", "REGISTRY_DELETION": "registry_delete",
    "PROCESS_MODULE_LOAD": "image_load", "PROCESS_OPEN": "process_access", "PROCESS_INJECTION": "create_remote_thread"}


WINDOWS_LOG_RE = re.compile(r"WINEVTLOG.*|WINDOWS.*|POWERSHELL.*|.*SYSMON.*", re.I)


NON_WINDOWS_ENDPOINT_RE = re.compile(r".*(LINUX|AUDITD|MAC|OSX|UNIX).*", re.I)


CONTEXT_FIELDS = {"metadata.event_type", "metadata.log_type", "metadata.product_event_type", "metadata.vendor_name",
                  "metadata.product_name", "metadata.base_labels.log_types"}   # scoping, not detection logic


def _dcc_dictionary(d: dict) -> list[dict]:
    """The installed calculator's scoring dictionary, read by column name (same aliases the calculator accepts)."""
    from openpyxl import load_workbook
    if not d["scoring"]:
        die(f"scoring dictionary not found next to {d['script']} - MITRE changed the layout; run yadda robustness --update")
    wb = load_workbook(d["scoring"], read_only=True, data_only=True)
    ws = next((w for w in wb.worksheets if {"Normalized Field Name", "Robustness Score"} <= {str(c or "").strip() for c in next(w.iter_rows(max_row=1, values_only=True))}), None)
    if ws is None:
        die(f"{d['scoring'].name}: no sheet with 'Normalized Field Name' and 'Robustness Score' - calculator format changed")
    rows = list(ws.iter_rows(values_only=True))
    head = [str(c or "").strip() for c in rows[0]]
    col = lambda *names: next((head.index(n) for n in names if n in head), None)
    ix = {"src": col("Log Source"), "eid": col("EventID", "Event ID"), "cat": col("Sigma Category", "Sigma logsource category"),
          "svc": col("Sigma Service", "Sigma logsource service"), "field": col("Normalized Field Name"),
          "ocsf": col("OCSF Schema Mapping", "OCSF Mapping"), "dc": col("Data Component")}
    if ix["ocsf"] is None:
        die(f"{d['scoring'].name}: no OCSF mapping column - the UDM->OCSF translation can't resolve fields")
    tok = lambda v: {x.strip().lower() for x in re.split(r"[;,]", str(v or "")) if x.strip()}
    out = []
    for r in rows[1:]:
        get = lambda k: r[ix[k]] if ix[k] is not None and ix[k] < len(r) else None
        if not get("field"):
            continue
        eid = str(get("eid") or "").strip()
        out.append({"src": str(get("src") or "").strip().lower(), "eid": re.sub(r"\.0+$", "", eid),
                    "cat": tok(get("cat")), "svc": tok(get("svc")), "field": str(get("field")).strip(),
                    "ocsf": re.sub(r"\s*\((derived|extension)\)", "", str(get("ocsf") or "")).strip().lower(),
                    "dc": str(get("dc") or "").strip()})
    return out


def _udm_ocsf() -> dict:
    out = defaultdict(list)
    f = SHARED / "udm_to_ocsf.csv"
    for r in csv.DictReader(f.open(encoding="utf-8")) if f.exists() else []:
        out[(r["udm_event_type"].strip() or "*", r["udm_field"].strip())].append(
            ([re.sub(r"\s*\((derived|extension)\)", "", o).strip().lower() for o in r["ocsf_fields"].split(";") if o.strip()],
             (r.get("sigma_field") or "").strip()))
    return out


def _shadow(rule: dict, text: str, dic: list[dict], u2o: dict, mapping: list[dict]) -> list[dict]:
    """One enabled YARA-L rule -> shadow Sigma documents. The events: section is parsed into its real boolean tree
    and kept in the Sigma condition (the calculator parses and/or/not/brackets). Per event variable, the tree is
    specialised for each event type it can match (e.g. PROCESS_LAUNCH or REGISTRY_*): those are OR-ed branches.
    Variables are AND-ed. Non-endpoint branches can't be scored for robustness; they carry their ATT&CK data
    component so the calculator's implementation coverage still works."""
    tree = _yl_events_tree(_yl_events(text))
    variables = sorted({p["var"] for p in _tree_preds(tree)})
    win_eids = {r["eid"] for r in dic if r["eid"]}
    docs = []
    for var in variables:
        vt = _tree_map(tree, lambda p: ("pred", p) if p["var"] == var else None)
        preds = _tree_preds(vt)
        if preds and all(p["field"].startswith("graph.") for p in preds):
            continue                                            # entity-graph lookup: context, not telemetry
        ets = sorted({x.upper() for p in preds if p["field"] == "metadata.event_type" and not p["neg"] for x in _values(p)})

        def only(p, et):                                        # the tree as it applies to one event type
            if p["field"] != "metadata.event_type" or not et or p.get("list"):
                return ("pred", p)
            return (et in {x.upper() for x in _values(p)}) != p["neg"]
        for et in ets or [""]:
            bt = _tree_map(vt, lambda p, et=et: only(p, et))
            if bt is False:
                continue
            bp = _tree_preds(bt)
            v = {"lt": {x.strip("^$() ") for p in bp if p["field"] == "metadata.log_type" and not p["neg"]
                        for x in p["value"].split("|") if x.strip("^$() ")},
                 "eid": {x for p in bp if p["field"] == "metadata.product_event_type" and not p["neg"] for x in re.findall(r"\d+", p["value"])},
                 "prod": " ".join(p["value"] for p in bp if p["field"] == "metadata.product_name")}
            docs.append(_shadow_doc(rule, var, et[:4], et, v, bt if bt is not True else None, dic, u2o, mapping, win_eids))
    return docs


ENDPOINT_PREFIX = re.compile(r"(PROCESS|FILE|REGISTRY|NETWORK_CONNECTION|NETWORK_DNS|SERVICE|SCHEDULED_TASK)_")


def _shadow_doc(rule, var, branch, et, v, bt, dic, u2o, mapping, win_eids) -> dict:
    windows_log = any(WINDOWS_LOG_RE.fullmatch(x) for x in v["lt"]) or bool(
        # no log type in the rule: only Windows when its product says so - "24" is also a OneLogin event code
        not v["lt"] and v["eid"] and re.search(r"windows|microsoft-windows|sysmon|service control manager|powershell|security-auditing", v["prod"], re.I))
    other_log = v["lt"] and not any(WINDOWS_LOG_RE.fullmatch(x) for x in v["lt"])
    non_win = bool(other_log and any(NON_WINDOWS_ENDPOINT_RE.fullmatch(x) for x in v["lt"]))
    endpoint = et in ENDPOINT_CATEGORIES and not non_win
    lts = f" ({', '.join(sorted(v['lt']))})" if v["lt"] else ""
    doc = {"rule": rule["name"], "var": var, "branch": branch, "techniques": rule["techniques"], "event_type": et,
           "unscored": [], "dropped": []}
    tags = [f"attack.{t.lower()}" for t in rule["techniques"]]
    if windows_log or endpoint:
        src = "sysmon" if any("SYSMON" in x.upper() for x in v["lt"]) else ("winlog" if windows_log and v["eid"] else "")
        ctx = [r for r in dic if (not v["eid"] or r["eid"] in v["eid"]) and (not src or r["src"].startswith(src[:3]))]
        cat = ENDPOINT_CATEGORIES.get(et, "")
        if not v["eid"]:
            ctx = [r for r in ctx if cat in r["cat"]] or ctx
        ctx.sort(key=lambda r: r["src"] != "sysmon")            # Sysmon rows first when the rule names no event ID
        logsource = {"product": "windows"}
        if v["eid"]:
            if ctx and all(r["cat"] for r in ctx):
                logsource["category"] = Counter(sorted(r["cat"])[0] for r in ctx).most_common(1)[0][0]
            if ctx and all(r["svc"] for r in ctx):
                logsource["service"] = Counter(sorted(r["svc"])[0] for r in ctx).most_common(1)[0][0]
        elif cat and any(cat in r["cat"] for r in dic):
            logsource["category"] = cat
        doc["basis"] = ("Windows event log" if windows_log else "endpoint, scored as Sysmon") + lts
        if v["eid"] and not ctx:
            doc["dropped"].append(f"event {', '.join(sorted(v['eid']))} not in the calculator's dictionary")

        def sigma_name(field):
            for ocsfs, override in u2o.get((et, field), []) + u2o.get(("*", field), []):
                if override and any(r["field"].lower() == override.lower() for r in ctx):
                    return next(r["field"] for r in ctx if r["field"].lower() == override.lower())
                for o in ocsfs:
                    hit = next((r["field"] for r in ctx if r["ocsf"] == o), None)
                    if hit:
                        return hit
            return None
        blocks, n = {}, [0]

        def leaf(p):
            if p["field"] in CONTEXT_FIELDS or p["field"].startswith("graph."):
                return None
            name = sigma_name(p["field"])
            if not name:
                doc["unscored"].append(p["field"])
                return None
            n[0] += 1
            key = f"{'filter' if p['neg'] else 'selection'}_{n[0]}"
            blocks[key] = {name + ("|re" if p["regex"] else ""): [p["value"]]}
            return ("ref", key, p["neg"])
        st = _tree_map(bt, leaf) if bt else None

        def cond(node, inside_not=False):
            if node[0] == "ref":
                if inside_not and node[1].startswith("selection_"):   # leaves under NOT(...) are exclusions
                    blocks[node[1].replace("selection_", "filter_")] = blocks.pop(node[1])
                    node = ("ref", node[1].replace("selection_", "filter_"), node[2])
                return ("not " if node[2] else "") + node[1]
            if node[0] == "not":
                return "not (" + cond(node[1], True) + ")"
            return "(" + f" {node[0]} ".join(cond(k, inside_not) for k in node[1]) + ")"
        detection = {}
        if v["eid"]:
            # present so the calculator picks the right dictionary rows; kept out of the condition, where it
            # would try to score "EventID" as a detection field
            detection["selection_eventid"] = {"EventID": [int(x) for x in sorted(v["eid"])]}
        if st and not isinstance(st, bool):
            text_cond = cond(st)
            detection.update(blocks)
            detection["condition"] = text_cond
            if all(k.startswith("filter_") for k in blocks):
                doc["dropped"].append("only exclusions left to score")
        else:
            doc["dropped"].append("no scorable field" + (" (only the event ID / fields the calculator doesn't know)" if v["eid"] or doc["unscored"] else ""))
            detection["selection_none"] = {"_no_scorable_field": "x"}
            detection["condition"] = "selection_none"
        doc["sigma"] = {"title": f"{rule['name']} [${var}]", "logsource": logsource, "tags": tags, "detection": detection}
        doc["kind"] = "scored"
        return doc
    # Not Windows/endpoint: no robustness scores exist for it. The implementation catalogue still covers every
    # platform, matched on ATT&CK data component - passed as the logsource category (the calculator's fallback).
    dcs = set()
    for lt in v["lt"] or {""}:
        for eid in v["eid"] or {""}:
            dcs |= _row_dcs({"event_type": et, "log_type": lt, "product_event_type": eid}, mapping)
    fields = sorted({p["field"] for p in _tree_preds(bt) if p["field"] not in CONTEXT_FIELDS and not p["field"].startswith("graph.")})
    windows_endpoint = bool(ENDPOINT_PREFIX.match(et)) and not other_log
    doc.update(basis=("Windows endpoint event the calculator has no Sigma category for" if windows_endpoint
                      else "not scorable - calculator has Windows/Sysmon scores only") + lts,
               kind="implementations only", dcs=sorted(dcs), unscored=fields)
    if not dcs:
        doc["dropped"].append("no ATT&CK data component known for its event/log type")
    base = {"title": f"{rule['name']} [${var}]", "tags": tags}
    product = sorted(v["lt"])[0].lower() if v["lt"] else ("windows" if windows_endpoint else "unknown")
    doc["sigma_multi"] = [dict(base, logsource={"product": product, "category": dc.lower().replace(" ", "_")},
                               detection={"selection": {f: ["x"] for f in fields[:20]} or {"_none": "x"}, "condition": "selection"})
                          for dc in sorted(dcs)] or [dict(base, logsource={"product": product, "category": "unknown"},
                                                          detection={"selection": {"_none": "x"}, "condition": "selection"})]
    doc["sigma"] = doc["sigma_multi"][0]
    return doc


def _dcc_sheet(wb, need: set, label: str):
    for ws in wb.worksheets:
        head = [str(c or "").strip() for c in next(ws.iter_rows(max_row=1, values_only=True), ())]
        if need <= set(head):
            return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(x is not None for x in r)]
    die(f"calculator output has no '{label}' sheet with columns {sorted(need)} - its output format changed; "
        f"the raw workbook is still in the environment's robustness folder")


def cmd_robustness(args):
    """yadda robustness <environment> [--update]  - MITRE Detection Coverage Calculator on the enabled YARA-L rules"""
    if "--update" in args:
        update_dcc()
        args = [a for a in args if a != "--update"]
        if not args:
            return
    if not args:
        die("usage: yadda robustness <environment> [--update]")
    c = envdir(args[0])
    d = _dcc()
    dic, u2o, mapping = _dcc_dictionary(d), _udm_ocsf(), _mapping()
    facts = _rule_facts(c)
    work = c / "robustness"
    shadow = work / "shadow"
    if shadow.exists():
        _rmtree(shadow)
    shadow.mkdir(parents=True)
    docs = []
    for r in facts["rules"]:
        f = c / "rules" / f"{r['name']}.yaral"
        if not f.exists():
            continue
        for i, doc in enumerate(_shadow(r, f.read_text(encoding="utf-8"), dic, u2o, mapping)):
            doc["files"] = []
            for j, sig in enumerate(doc.get("sigma_multi") or [doc["sigma"]]):
                fn = f"{r['name'][:60]}_{hashlib.sha1(r['name'].encode()).hexdigest()[:6]}__{doc['var'][:20]}__{i}_{j}.yml"
                write_text(shadow / fn, yaml_dump_plain(sig))
                doc["files"].append(fn)
            docs.append(doc)
    if not docs:
        print(f"{c.name}: no enabled rules yet - nothing for MITRE's calculator to score")
        return
    print(f"{c.name}: {len({x['rule'] for x in docs})} enabled rules -> {len(docs)} shadow Sigma files; "
          f"running MITRE Detection Coverage Calculator ({d['version']}) ...", flush=True)
    out_prefix = work / "dcc_results"
    p = subprocess.run([sys.executable, str(d["script"]), str(shadow), "--out-prefix", str(out_prefix)],
                       cwd=d["script"].parent, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if p.returncode != 0:
        print(p.stdout[-2000:], p.stderr[-3000:])
        die("the calculator failed (output above). If MITRE changed its command line, check: "
            f"{sys.executable} {d['script']} --help")
    from openpyxl import load_workbook
    wb = load_workbook(out_prefix.with_suffix(".xlsx"), read_only=True, data_only=True)
    scores = _dcc_sheet(wb, {"Analytic File", "Robustness Score", "Precision Score"}, "Analytic Scores")
    techcov = _dcc_sheet(wb, {"ATT&CK Technique ID", "Matched Implementation Count", "Total Implementation Count"}, "Technique Coverage")
    num = lambda x: int(float(x)) if str(x or "").strip() not in ("", "None") else None
    field_rows = defaultdict(list)          # optional sheet: which field scored what (explains a rule's score)
    for ws in wb.worksheets:
        head = [str(x or "").strip() for x in next(ws.iter_rows(max_row=1, values_only=True), ())]
        if {"Analytic File", "Analytic Field", "Robustness Score", "Matched"} <= set(head):
            for row in ws.iter_rows(min_row=2, values_only=True):
                rd = dict(zip(head, row))
                if str(rd.get("Matched") or "").lower() == "yes" and num(rd.get("Robustness Score")) is not None:
                    field_rows[str(rd["Analytic File"])].append((str(rd["Analytic Field"]), num(rd["Robustness Score"]),
                                                                  str(rd.get("Field Role") or "")))
    by_file = {}
    for row in scores:
        by_file.setdefault(str(row.get("Analytic File") or ""), row)
    # Combine like the calculator does: OR-ed branches of one event variable -> strongest robustness, weakest
    # precision; AND-ed event variables -> weakest robustness, strongest precision. Unscored parts are flagged.
    for doc in docs:
        rows_ = [by_file.get(fn, {}) for fn in doc["files"]]
        doc["robustness"] = num(rows_[0].get("Robustness Score")) if doc["kind"] == "scored" and rows_ else None
        doc["precision"] = num(rows_[0].get("Precision Score")) if doc["kind"] == "scored" and rows_ else None
        doc["missing"] = str(rows_[0].get("Missing Scores") or "") if rows_ else ""
    per_rule = defaultdict(lambda: defaultdict(list))
    for doc in docs:
        per_rule[doc["rule"]][doc["var"]].append(doc)
    rule_rows = []
    for name, vars_ in sorted(per_rule.items()):
        alld = [x for ds in vars_.values() for x in ds]
        var_scores, gaps = [], []
        for var, ds in vars_.items():
            sc = [x for x in ds if x["kind"] == "scored" and x["robustness"] is not None]
            if sc:
                best = max(x["robustness"] for x in sc)
                var_scores.append((best, min(x["precision"] for x in sc if x["robustness"] == best)))
            if len(sc) < len(ds):
                for x in ds:
                    if x not in sc:
                        gaps.append(f"${var}: " + ("; ".join(x["dropped"]) or (x["basis"] if x["kind"] != "scored" else f"unscored field(s) {x['missing'] or ', '.join(x['unscored'])}")))
        scored_any = any(x["kind"] == "scored" for x in alld)
        rob = min(v[0] for v in var_scores) if var_scores else ""
        prec = max(v[1] for v in var_scores if v[0] == rob) if var_scores else ""
        if var_scores:
            status = "scored" if not gaps and not any(x["unscored"] for x in alld if x["kind"] == "scored") else "scored (partial)"
        else:
            status = "not scored" if scored_any else "implementations only (not Windows)"
        weakest = sorted({f"{f} ({'exclusion' if 'filter' in role else 'match'})"
                          for x in alld for fn in x["files"] for f, sc, role in field_rows.get(fn, []) if rob != "" and sc == rob})
        rule_rows.append({"rule": name, "techniques": " ".join(alld[0]["techniques"]), "robustness": rob, "precision": prec,
                          "weakest_fields": "; ".join(weakest),
                          "status": status, "basis": "; ".join(sorted({x["basis"] for x in alld})),
                          "why_not_fully_scored": " | ".join(dict.fromkeys(gaps))[:400],
                          "fields_not_scored": "; ".join(sorted({f for x in alld if x["kind"] == "scored" for f in x["unscored"]})),
                          "event_variables": len(vars_), "branches": len(alld),
                          "shadow_files": " ".join(fn for x in alld for fn in x["files"])})
    tech_rows = []
    rules_by_tech = defaultdict(list)
    for r in rule_rows:
        for t in r["techniques"].split():
            rules_by_tech[t].append(r)
    for row in techcov:
        tid = str(row.get("ATT&CK Technique ID") or "").strip()
        tot, m = num(row.get("Total Implementation Count")) or 0, num(row.get("Matched Implementation Count")) or 0
        rs = rules_by_tech.get(tid, [])
        robs = [r["robustness"] for r in rs if r["robustness"] != ""]
        # several rules on one technique are OR-ed: an attacker has to evade all of them -> strongest robustness
        tech_rows.append({"technique": tid, "implementations_matched": m, "implementations_total": tot,
                          "implementation_coverage": round(m / tot, 3) if tot else "",
                          "best_robustness": max(robs) if robs else "",
                          "precision_of_best": max((r["precision"] for r in rs if r["robustness"] == max(robs)), default="") if robs else "",
                          "rules": " ".join(r["rule"] for r in rs),
                          "scored_rules": sum(1 for r in rs if r["status"].startswith("scored")),
                          "calculator": d["version"]})
    for name, rows_, fields in (("robustness_rules.csv", rule_rows, list(rule_rows[0])),
                                ("robustness_techniques.csv", tech_rows, ["technique", "implementations_matched", "implementations_total",
                                                                         "implementation_coverage", "best_robustness", "precision_of_best", "rules", "scored_rules", "calculator"])):
        write_csv(c / name, fields, rows_)
    st = Counter(r["status"] for r in rule_rows)
    band = Counter(r["robustness"] for r in rule_rows if r["robustness"] != "")
    unmapped = Counter(f for r in rule_rows for f in r["fields_not_scored"].split("; ") if f)
    print(f"  rules: {st['scored']} fully scored, {st['scored (partial)']} partly scored, {st['not scored']} not scorable "
          f"(event or fields not in the calculator), {st['implementations only (not Windows)']} non-Windows (implementation coverage only)")
    print("  robustness (1 = ephemeral value ... 5 = core to the technique): " + (", ".join(f"{k}: {band[k]} rules" for k in sorted(band)) or "none scored"))
    print(f"  techniques: {sum(1 for t in tech_rows if t['implementations_total'])} of {len(tech_rows)} tagged techniques are in MITRE's implementation catalogue")
    if unmapped:
        print("  UDM fields the calculator has no field for (add to shared/udm_to_ocsf.csv if one exists): "
              + ", ".join(f"{k} ({v})" for k, v in unmapped.most_common(8)))
    print(f"  -> {c / 'robustness_rules.csv'}\n  -> {c / 'robustness_techniques.csv'}\n  -> calculator workbook: {out_prefix.with_suffix('.xlsx')}")


def _robustness_rules(c: Path) -> dict:
    f = c / "robustness_rules.csv"
    return {r["rule"]: r for r in csv.DictReader(f.open(encoding="utf-8-sig"))} if f.exists() else {}


def _robustness(c: Path) -> dict:
    """technique -> calculator results (empty if yadda robustness hasn't been run)."""
    f = c / "robustness_techniques.csv"
    out = {}
    for r in read_csv(f)[1] if f.exists() else []:
        tot, matched = int(r["implementations_total"] or 0), int(r["implementations_matched"] or 0)
        # 0 matched because the calculator can't read cloud/SaaS rules is "unknown", not "0% deep"
        readable = matched > 0 or int(r.get("scored_rules") or 0) > 0
        out[r["technique"]] = {"matched": matched, "total": tot,
                               "impl": (matched / tot) if tot and readable else None,
                               "robustness": int(r["best_robustness"]) if r["best_robustness"] else None,
                               "precision": int(r["precision_of_best"]) if r["precision_of_best"] else None}
    return out
