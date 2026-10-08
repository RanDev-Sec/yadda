"""telemetry - reads the SecOps telemetry inventory export and maps log types / UDM event types to ATT&CK data
components (shared/udm_event_type_to_data_component.csv). `yadda mapping` checks that table against the Windows and
Sysmon mappings shipped with MITRE's Detection Coverage Calculator."""
from __future__ import annotations

import datetime as dt
import json
import re
import shutil
from collections import Counter, defaultdict
from pathlib import Path
from delib.config import ENVIRONMENTS, SHARED, STIX, die, read_csv, write_csv
from delib.attack import EXTERNAL_DCS, attack_version
from delib.upstream import _dcc


def _read_inventory(path: Path) -> list[dict]:
    """Read a SecOps export of telemetry_inventory.yaral, tolerating what Excel and SecOps do to it:
    '$event_count' / 'event_count' / 'events', thousands separators, BOM, cp1252, semicolons, footer rows."""
    head, raw = read_csv(path)
    norm = {h: h.strip().lstrip("$").lower() for h in head if h}
    def col(*names):
        return next((h for h, n in norm.items() if n in names), None)
    lt, et, pet = col("log_type"), col("event_type"), col("product_event_type")
    ev = col("event_count", "events", "count")
    if not (lt and et and ev):
        die(f"{path.name}: expected columns log_type, event_type, event_count - found {list(norm.values())}")
    rows = []
    for r in raw:
        log_type, event_type = (r.get(lt) or "").strip(), (r.get(et) or "").strip()
        if not log_type or not event_type or log_type.lower() in ("total", "sum", "grand total"):
            continue                                            # footer / subtotal / half-empty export rows
        rows.append({"log_type": log_type, "event_type": event_type, "events": _count(r.get(ev)),
                     "product_event_type": (r.get(pet) or "").strip() if pet else ""})
    return rows


def _read_export(path: Path, columns: dict) -> list[dict]:
    """Read a SecOps CSV export. columns = {key: (accepted header names...)}; headers may carry '$'."""
    head, raw = read_csv(path)
    norm = {h: h.strip().lstrip("$").lower() for h in head if h}
    found = {k: next((h for h, n in norm.items() if n in names), None) for k, names in columns.items()}
    missing = [k for k, h in found.items() if h is None]
    if missing:
        die(f"{path.name}: missing column(s) {missing} - found {list(norm.values())}")
    return [{k: (r.get(h) or "").strip() for k, h in found.items()} for r in raw]


def present(inventory: list[dict], min_events: int = 100) -> dict:
    """Which event types / log types count as received: 100+ events in the export window (summed over rows), or
    any row carrying a specific product event ID (a rare event ID still proves the source).
    {'event_types': set, 'log_types': set, 'rows': [rows that count]}"""
    et, lt = Counter(), Counter()
    for r in inventory:
        et[r["event_type"]] += r["events"]
        lt[r["log_type"]] += r["events"]
    keep = [r for r in inventory if (r["product_event_type"] and r["events"] > 0)
            or et[r["event_type"]] >= min_events or r["events"] >= min_events]
    return {"event_types": {r["event_type"] for r in keep}, "log_types": {r["log_type"] for r in keep}, "rows": keep}


def _clean_number(v) -> str:
    return str(v if v is not None else "").strip().strip('"\'').replace("\u00a0", "").replace(" ", "").replace(",", "")


def _num(v) -> float:
    """'1,234' / '1 234' / '"12"' / '1.234.567' -> number; anything unreadable ('N/A', '') -> 0."""
    v = _clean_number(v)
    if re.fullmatch(r"\d{1,3}(\.\d{3}){2,}", v):            # 1.234.567: dots as thousands separators
        v = v.replace(".", "")
    try:
        return float(v) if v else 0.0
    except ValueError:
        return 0.0


def _count(v) -> int:
    """A whole-number count; '1.234' here can only mean a thousands separator (European export)."""
    s = _clean_number(v)
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    return int(_num(s))


_DATE_FORMATS = ("%Y/%m/%d", "%b %d, %Y", "%B %d, %Y", "%d %b %Y", "%d %B %Y", "%b %d %Y", "%d-%b-%Y", "%Y%m%d")


def day_first(values) -> bool:
    """True when a column of n/n/yyyy dates is day-first: some first number is over 12 and none is a month-only
    value (second number over 12). Decided once per file, so 05/06/2025 is read the same way as its neighbours."""
    firsts, seconds = [], []
    for v in values:
        m = re.match(r"\s*\"?(\d{1,2})[/.-](\d{1,2})[/.-]\d{4}\b", str(v or ""))
        if m:
            firsts.append(int(m.group(1)))
            seconds.append(int(m.group(2)))
    return any(a > 12 for a in firsts) and not any(b > 12 for b in seconds)


def _when(v, dayfirst: bool | None = None):
    """A last-detection time from an export -> date, or None for never/0/empty/unreadable.
    Accepts epoch seconds / milliseconds / microseconds, ISO 8601 (also 2025-8-31), 2025/08/31, n/n/yyyy
    (`dayfirst` from day_first() over the whole column; without it, US month-first unless the first number is
    over 12), 'Aug 31, 2025' / 'Sept 1, 2025' and '31 Aug 2025', each optionally followed by a time."""
    v = str(v if v is not None else "").strip().strip('"')
    if not v or v.lower() in ("0", "0.0", "never", "none", "null", "n/a", "-"):
        return None
    if re.fullmatch(r"\d+(\.\d+)?", v.replace(",", "")) and len(v.split(".")[0]) >= 9:
        digits = len(v.replace(",", "").split(".")[0])         # 10 = seconds, 13 = ms, 16 = µs, 19 = ns
        if digits > 20:
            return None
        secs = _num(v) / 1000 ** max(0, (digits - 10 + 1) // 3)
        try:
            return dt.datetime.fromtimestamp(secs, dt.timezone.utc).date() if secs > 0 else None
        except (ValueError, OverflowError, OSError):
            return None
    m = re.match(r"(\d{4})[/.-](\d{1,2})[/.-](\d{1,2})(?!\d)", v)
    if m:
        try:
            return dt.date(*map(int, m.groups()))
        except ValueError:
            return None
    m = re.match(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})(?!\d)", v)
    if m:
        a, b, y = map(int, m.groups())
        first_is_day = dayfirst if dayfirst is not None else a > 12
        month, day = (b, a) if first_is_day else (a, b)
        try:
            return dt.date(y, month, day)
        except ValueError:
            return None
    head = re.split(r"[ T,]\s*(?=\d{1,2}:\d{2})", v)[0].strip(" ,")
    head = re.sub(r"\bSept\b", "Sep", head, flags=re.I)
    for fmt in _DATE_FORMATS:
        try:
            return dt.datetime.strptime(head, fmt).date()
        except ValueError:
            continue
    return None


def _mapping() -> list[dict]:
    """shared/udm_event_type_to_data_component.csv. A row matches an inventory row when its event type matches
    ('*' = any) and, if given, its log_type / product_event_type regexes fully match (case-insensitive)."""
    out = []
    for r in read_csv(SHARED / "udm_event_type_to_data_component.csv")[1]:
        lt, pet = (r.get("log_type") or "").strip(), (r.get("product_event_type") or "").strip()
        out.append({"et": r["udm_event_type"].strip(), "dc": r["data_component"].strip(),
                    "lt": re.compile(lt, re.I) if lt else None, "pet": re.compile(pet, re.I) if pet else None,
                    "label": _pretty(lt, pet, r["udm_event_type"].strip())})
    return out


def _pretty(lt: str, pet: str, et: str) -> str:
    """Readable label for a mapping row: 'WINEVTLOG event 4663', 'AWS_CLOUDTRAIL RunInstances', 'PROCESS_LAUNCH'."""
    lt = re.split(r"\|", lt)[0].replace(".*", "") if lt else ""
    m = re.fullmatch(r"\(([\d|]+)\)\(\\D\.\*\)\?", pet or "")
    if m:
        pet = "event " + m.group(1).replace("|", "/")
    elif pet:
        pet = re.sub(r"\\\.", ".", pet).replace(".*", "").replace("(", "").replace(")", "").replace("|", "/")
    return " ".join(x for x in (lt, pet or et) if x)


def _row_dcs(row: dict, mapping: list[dict]) -> set:
    out = set()
    for m in mapping:
        if m["et"] not in ("*", row["event_type"]):
            continue
        if m["lt"] and not m["lt"].fullmatch(row["log_type"]):
            continue
        if m["pet"] and not (row.get("product_event_type") and m["pet"].fullmatch(row["product_event_type"])):
            continue
        out.add(m["dc"])
    return out


def _dc_sources(mapping: list[dict]) -> dict:
    out = defaultdict(list)
    for m in mapping:
        if m["label"] not in out[m["dc"]]:
            out[m["dc"]].append(m["label"])
    return out


# ---------------------------------------------------------------- mapping check against MITRE's DCC mappings.xlsx
DCC_LOGTYPE = {"winlog": "WINEVTLOG.*|POWERSHELL.*", "sysmon": "WINDOWS_SYSMON|SYSMON.*"}


DCC_SAMPLE_LT = {"winlog": "WINEVTLOG", "sysmon": "WINDOWS_SYSMON"}


def _dcc_mappings(d: dict) -> dict:
    """(winlog|sysmon, event id) -> {data component: message summary} from the installed calculator's mappings.xlsx."""
    from openpyxl import load_workbook
    if not d["mappings"]:
        die(f"mappings workbook not found next to {d['script']} - run yadda robustness --update")
    wb = load_workbook(d["mappings"], read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        head = [str(c or "").strip() for c in rows[0]] if rows else []
        dc_col = next((i for i, h in enumerate(head) if "data component" in h.lower()), None)
        if {"Log Source", "EventID"} <= set(head) and dc_col is not None:
            out = defaultdict(dict)
            for r in rows[1:]:
                r = list(r) + [None] * (len(head) - len(r))
                src, eid, dc = str(r[head.index("Log Source")] or "").lower(), str(r[head.index("EventID")] or "").strip(), str(r[dc_col] or "").strip()
                kind = "sysmon" if "sysmon" in src else "winlog" if "windows" in src or "winlog" in src else None
                eid = re.sub(r"\.0+$", "", eid)
                if kind and eid.isdigit() and dc:
                    summ = str(r[head.index("Message Summary")] or "") if "Message Summary" in head else ""
                    out[(kind, eid)].setdefault(dc, summ)
            return out
    die(f"{d['mappings'].name}: no sheet with Log Source / EventID / data component columns - calculator format changed")


def cmd_mapping(args):
    """yadda mapping [--apply]  - check shared/udm_event_type_to_data_component.csv against MITRE's Windows/Sysmon mappings"""
    d = _dcc()
    dcc = _dcc_mappings(d)
    mapping = _mapping()
    known_dcs = {o["name"] for o in json.load((STIX / "enterprise-attack" / "enterprise-attack.json").open(encoding="utf-8"))["objects"]
                 if o["type"] == "x-mitre-data-component" and not o.get("revoked") and not o.get("x_mitre_deprecated")}
    # SecOps event type actually seen for each (log type family, event id), from every environment's inventory
    seen = defaultdict(set)
    from delib import inputs
    invs = [f for c in ENVIRONMENTS.iterdir() if c.is_dir() for f in [inputs.current(c, "inventory")] if f]
    for inv in invs:
        for r in _read_inventory(inv):
            if r["product_event_type"]:
                kind = "sysmon" if re.fullmatch(DCC_LOGTYPE["sysmon"], r["log_type"], re.I) else \
                    "winlog" if re.fullmatch(DCC_LOGTYPE["winlog"], r["log_type"], re.I) else None
                for eid in re.findall(r"\d+", r["product_event_type"])[:1]:
                    if kind:
                        seen[(kind, eid)].add((r["log_type"], r["event_type"]))
    report, add = [], defaultdict(list)
    for (kind, eid), dcs in sorted(dcc.items(), key=lambda kv: (kv[0][0], int(kv[0][1]))):
        samples = seen.get((kind, eid)) or {(DCC_SAMPLE_LT[kind], "")}
        ours = set().union(*(_row_dcs({"event_type": et, "log_type": lt, "product_event_type": eid}, mapping) for lt, et in samples))
        basis = "inventory event type " + "/".join(sorted({et for _, et in samples if et})) if seen.get((kind, eid)) else "explicit rows only"
        for dc, summ in dcs.items():
            if dc.lower() in ("no equivalent component", "none"):
                continue
            if dc not in known_dcs:
                report.append({"source": kind, "event_id": eid, "data_component": dc, "status": "MITRE name not in ATT&CK " + attack_version(), "basis": basis, "summary": summ})
            elif dc in EXTERNAL_DCS:
                report.append({"source": kind, "event_id": eid, "data_component": dc, "status": "skipped - external (not environment telemetry)", "basis": basis, "summary": summ})
            elif dc in ours:
                report.append({"source": kind, "event_id": eid, "data_component": dc, "status": "agree", "basis": basis, "summary": summ})
            else:
                report.append({"source": kind, "event_id": eid, "data_component": dc, "status": "missing - add", "basis": basis, "summary": summ})
                add[(kind, dc)].append((eid, summ))
        explicit = {m["dc"] for m in mapping if m["pet"] and m["lt"] and m["lt"].fullmatch(DCC_SAMPLE_LT[kind]) and m["pet"].fullmatch(eid)}
        for dc in sorted(explicit - set(dcs)):
            report.append({"source": kind, "event_id": eid, "data_component": dc, "status": "ours only - review",
                           "basis": "explicit row", "summary": "MITRE maps this event to: " + ", ".join(sorted(dcs))})
    out = SHARED / "mapping_check.csv"
    write_csv(out, ["status", "source", "event_id", "data_component", "basis", "summary"],
              sorted(report, key=lambda r: (r["status"], r["source"], int(r["event_id"]))))
    st = Counter(r["status"] for r in report)
    print(f"MITRE Detection Coverage Calculator mappings ({d['version']}): {len(dcc)} Windows/Sysmon event IDs")
    for k, v in st.most_common():
        print(f"  {k:45} {v}")
    newly = {dc for _, dc in add} - set(_dc_sources(mapping))
    print(f"  data components that become measurable: {len(newly)}" + (f" ({', '.join(sorted(newly))})" if newly else ""))
    for r in [r for r in report if r["status"] == "ours only - review"][:12]:
        print(f"  review: {r['source']} {r['event_id']}: we map {r['data_component']}; {r['summary']}")
    print(f"  full report: {out}")
    if "--apply" not in args:
        if add:
            print("  run  yadda mapping --apply  to add the missing rows (disagreements are never changed automatically)")
        return
    f = SHARED / "udm_event_type_to_data_component.csv"
    shutil.copy(f, f.with_name(f.stem + f".backup-{dt.datetime.now():%Y%m%d-%H%M%S}.csv"))
    head, rows = read_csv(f)
    fields = list(dict.fromkeys(head + ["udm_event_type", "data_component", "note", "log_type", "product_event_type"]))
    for (kind, dc), items in sorted(add.items()):
        eids = sorted({e for e, _ in items}, key=int)
        rows.append({"udm_event_type": "*", "data_component": dc,
                     "note": f"MITRE DCC mappings.xlsx ({d['version']}): " + "; ".join(sorted({f'{e} {s[:40]}' for e, s in items}))[:180],
                     "log_type": DCC_LOGTYPE[kind], "product_event_type": f"({'|'.join(eids)})(\\D.*)?"})
    write_csv(f, fields, rows, encoding="utf-8")
    print(f"  added {len(add)} rows to {f.name} (backup kept next to it). Re-run yadda data / yadda dashboard to use them.")
